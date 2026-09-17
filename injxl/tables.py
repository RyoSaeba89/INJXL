# -*- coding: utf-8 -*-
"""Lecture / ecriture des tableaux (xlsx, xlsm, csv, xls en lecture seule).

Toutes les coordonnees sont 1-base, comme dans Excel.
"""

from __future__ import annotations

import csv
import io
import os
import shutil
import datetime
from typing import Any, List, Optional

try:
    import openpyxl
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter, column_index_from_string
    OPENPYXL_OK = True
except ImportError:  # pragma: no cover
    OPENPYXL_OK = False
    get_column_letter = None
    column_index_from_string = None

EXT_EXCEL = (".xlsx", ".xlsm", ".xltx", ".xltm")
EXT_CSV = (".csv", ".txt", ".tsv")
EXT_XLS = (".xls",)

MAX_SCAN_VIDE = 5000   # nb de lignes vides consecutives tolerees en fin de feuille


def lettre_colonne(index: int) -> str:
    """1 -> A, 27 -> AA (sans dependre d'openpyxl)."""
    if index < 1:
        return "?"
    s = ""
    while index:
        index, reste = divmod(index - 1, 26)
        s = chr(65 + reste) + s
    return s


def index_colonne(lettre: str) -> int:
    """A -> 1, AA -> 27. Retourne 0 si invalide."""
    lettre = (lettre or "").strip().upper()
    if not lettre or not lettre.isalpha():
        return 0
    n = 0
    for c in lettre:
        n = n * 26 + (ord(c) - 64)
    return n


def type_fichier(path: str) -> str:
    ext = os.path.splitext(path or "")[1].lower()
    if ext in EXT_EXCEL:
        return "xlsx"
    if ext in EXT_CSV:
        return "csv"
    if ext in EXT_XLS:
        return "xls"
    return "inconnu"


def est_vide(valeur: Any) -> bool:
    return valeur is None or (isinstance(valeur, str) and valeur.strip() == "")


def lister_feuilles(path: str) -> List[str]:
    """Noms des feuilles d'un classeur ; pour un CSV, une pseudo-feuille."""
    kind = type_fichier(path)
    if kind == "csv":
        return ["(fichier texte)"]
    if kind == "xls":
        import xlrd  # leve ImportError si absent : message gere par l'appelant
        livre = xlrd.open_workbook(path, on_demand=True)
        noms = livre.sheet_names()
        livre.release_resources()
        return list(noms)
    if not OPENPYXL_OK:
        raise ImportError("openpyxl est requis pour lire les fichiers Excel.")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


# ---------------------------------------------------------------------------
#  Tableaux
# ---------------------------------------------------------------------------

class Table(object):
    """Interface commune. Coordonnees 1-base."""

    def __init__(self, path: str, sheet: str, header_row: int):
        self.path = path
        self.sheet = sheet
        self.header_row = max(1, int(header_row or 1))
        self.max_row = 0
        self.max_column = 0
        self.writable = False
        self._modifiees = set()

    # --- lecture ---
    def get(self, row: int, col: int) -> Any:
        raise NotImplementedError

    def headers(self) -> List[str]:
        out = []
        for c in range(1, self.max_column + 1):
            v = self.get(self.header_row, c)
            out.append("" if v is None else str(v).strip())
        return out

    def nb_lignes_donnees(self) -> int:
        return max(0, self.max_row - self.header_row)

    def apercu(self, nb: int = 50) -> List[List[Any]]:
        lignes = []
        for r in range(self.header_row + 1, min(self.max_row, self.header_row + nb) + 1):
            lignes.append([self.get(r, c) for c in range(1, self.max_column + 1)])
        return lignes

    # --- ecriture ---
    def set(self, row: int, col: int, value: Any) -> None:
        raise NotImplementedError

    def ajouter_colonne(self, header: str) -> int:
        raise NotImplementedError

    def surligner(self, row: int, col: int, couleur: str) -> None:
        pass

    def save(self, path: str) -> None:
        raise NotImplementedError

    def close(self) -> None:
        pass


class ExcelTable(Table):
    """Classeur .xlsx / .xlsm.

    En lecture seule les cellules sont materialisees en memoire (rapide) ;
    en ecriture le classeur reste ouvert pour preserver formules et mise en forme.
    """

    def __init__(self, path: str, sheet: str = "", header_row: int = 1, writable: bool = False):
        Table.__init__(self, path, sheet, header_row)
        if not OPENPYXL_OK:
            raise ImportError("openpyxl est requis pour lire les fichiers Excel.")
        self.writable = writable
        self._rows = None
        if writable:
            self.wb = openpyxl.load_workbook(path, data_only=False, keep_vba=path.lower().endswith(".xlsm"))
            self.ws = self.wb[sheet] if sheet and sheet in self.wb.sheetnames else self.wb.active
            self.sheet = self.ws.title
            self.max_column = self.ws.max_column or 1
            self.max_row = self._dernier_non_vide(self.ws.max_row or 1)
        else:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            try:
                ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
                self.sheet = ws.title
                self._rows = [list(r) for r in ws.iter_rows(values_only=True)]
            finally:
                wb.close()
            while self._rows and all(est_vide(v) for v in self._rows[-1]):
                self._rows.pop()
            self.max_row = len(self._rows)
            self.max_column = max([len(r) for r in self._rows] or [1])
            self.wb = None
            self.ws = None

    def _dernier_non_vide(self, depart: int) -> int:
        """Remonte depuis la derniere ligne declaree jusqu'a une ligne non vide."""
        ligne = depart
        vides = 0
        while ligne > self.header_row and vides < MAX_SCAN_VIDE:
            if any(not est_vide(self.ws.cell(row=ligne, column=c).value)
                   for c in range(1, self.max_column + 1)):
                return ligne
            vides += 1
            ligne -= 1
        return max(ligne, self.header_row)

    def get(self, row, col):
        if self._rows is not None:
            if 1 <= row <= len(self._rows):
                ligne = self._rows[row - 1]
                if 1 <= col <= len(ligne):
                    return ligne[col - 1]
            return None
        return self.ws.cell(row=row, column=col).value

    def set(self, row, col, value):
        self.ws.cell(row=row, column=col).value = value
        self._modifiees.add((row, col))
        if row > self.max_row:
            self.max_row = row

    def ajouter_colonne(self, header):
        self.max_column += 1
        self.ws.cell(row=self.header_row, column=self.max_column).value = header
        try:
            modele = self.ws.cell(row=self.header_row, column=max(1, self.max_column - 1))
            cible = self.ws.cell(row=self.header_row, column=self.max_column)
            cible.font = modele.font.copy()
            cible.fill = modele.fill.copy()
            cible.border = modele.border.copy()
            cible.alignment = modele.alignment.copy()
        except Exception:
            pass
        return self.max_column

    def surligner(self, row, col, couleur):
        remplissage = PatternFill(start_color=couleur, end_color=couleur, fill_type="solid")
        self.ws.cell(row=row, column=col).fill = remplissage

    def save(self, path):
        self.wb.save(path)

    def close(self):
        try:
            if self.wb is not None:
                self.wb.close()
        except Exception:
            pass


class XlsTable(Table):
    """Ancien format .xls, en lecture seule (via xlrd)."""

    def __init__(self, path: str, sheet: str = "", header_row: int = 1, writable: bool = False):
        Table.__init__(self, path, sheet, header_row)
        if writable:
            raise ValueError("Le format .xls ne peut pas etre modifie par INJXL. "
                             "Enregistrez le fichier cible en .xlsx.")
        import xlrd
        livre = xlrd.open_workbook(path)
        feuille = livre.sheet_by_name(sheet) if sheet and sheet in livre.sheet_names() \
            else livre.sheet_by_index(0)
        self.sheet = feuille.name
        self._rows = [[feuille.cell_value(r, c) for c in range(feuille.ncols)]
                      for r in range(feuille.nrows)]
        while self._rows and all(est_vide(v) for v in self._rows[-1]):
            self._rows.pop()
        self.max_row = len(self._rows)
        self.max_column = max([len(r) for r in self._rows] or [1])

    def get(self, row, col):
        if 1 <= row <= len(self._rows):
            ligne = self._rows[row - 1]
            if 1 <= col <= len(ligne):
                v = ligne[col - 1]
                return None if v == "" else v
        return None


class CsvTable(Table):
    """Fichier texte delimite. Le delimiteur et l'encodage sont detectes."""

    def __init__(self, path: str, sheet: str = "", header_row: int = 1, writable: bool = False):
        Table.__init__(self, path, sheet, header_row)
        self.writable = writable
        self.encodage, texte = self._lire(path)
        self.delimiteur = self._detecter_delimiteur(texte)
        lecteur = csv.reader(io.StringIO(texte), delimiter=self.delimiteur)
        self._rows = [list(r) for r in lecteur]
        while self._rows and all(est_vide(v) for v in self._rows[-1]):
            self._rows.pop()
        self.max_row = len(self._rows)
        self.max_column = max([len(r) for r in self._rows] or [1])
        self.sheet = "(fichier texte)"

    @staticmethod
    def _lire(path):
        donnees = open(path, "rb").read()
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return enc, donnees.decode(enc)
            except UnicodeDecodeError:
                continue
        return "latin-1", donnees.decode("latin-1", "replace")

    @staticmethod
    def _detecter_delimiteur(texte):
        echantillon = texte[:8192]
        try:
            return csv.Sniffer().sniff(echantillon, delimiters=";,\t|").delimiter
        except Exception:
            comptes = {d: echantillon.count(d) for d in ";,\t|"}
            meilleur = max(comptes, key=comptes.get)
            return meilleur if comptes[meilleur] else ";"

    def get(self, row, col):
        if 1 <= row <= len(self._rows):
            ligne = self._rows[row - 1]
            if 1 <= col <= len(ligne):
                v = ligne[col - 1]
                return None if v == "" else v
        return None

    def _garantir(self, row, col):
        while len(self._rows) < row:
            self._rows.append([])
        ligne = self._rows[row - 1]
        while len(ligne) < col:
            ligne.append("")

    def set(self, row, col, value):
        self._garantir(row, col)
        self._rows[row - 1][col - 1] = "" if value is None else str(value)
        self._modifiees.add((row, col))
        self.max_row = max(self.max_row, row)
        self.max_column = max(self.max_column, col)

    def ajouter_colonne(self, header):
        self.max_column += 1
        self.set(self.header_row, self.max_column, header)
        return self.max_column

    def save(self, path):
        enc = "utf-8-sig" if self.encodage.startswith("utf-8") else self.encodage
        with open(path, "w", encoding=enc, newline="") as f:
            ecrivain = csv.writer(f, delimiter=self.delimiteur, quoting=csv.QUOTE_MINIMAL)
            for ligne in self._rows:
                ecrivain.writerow(ligne)


def ouvrir(path: str, sheet: str = "", header_row: int = 1, writable: bool = False) -> Table:
    """Fabrique : retourne le Table adapte a l'extension."""
    if not path or not os.path.isfile(path):
        raise FileNotFoundError("Fichier introuvable : %s" % path)
    kind = type_fichier(path)
    if kind == "xlsx":
        return ExcelTable(path, sheet, header_row, writable)
    if kind == "csv":
        return CsvTable(path, sheet, header_row, writable)
    if kind == "xls":
        return XlsTable(path, sheet, header_row, writable)
    raise ValueError("Format non pris en charge : %s" % os.path.basename(path))


def horodatage() -> str:
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def chemin_auto(path_cible: str, suffixe: str = "INJXL") -> str:
    """Nom de sortie horodate, dans le dossier du fichier cible."""
    dossier = os.path.dirname(os.path.abspath(path_cible))
    base, ext = os.path.splitext(os.path.basename(path_cible))
    return os.path.join(dossier, "%s_%s_%s%s" % (base, suffixe, horodatage(), ext or ".xlsx"))


def sauvegarder_copie(path: str) -> str:
    """Duplique un fichier avant modification sur place. Retourne le chemin de la copie."""
    base, ext = os.path.splitext(os.path.abspath(path))
    copie = "%s_SAUVEGARDE_%s%s" % (base, horodatage(), ext)
    shutil.copy2(path, copie)
    return copie
