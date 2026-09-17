# -*- coding: utf-8 -*-
"""Moteur d'INJXL : normalisation des cles, indexation, analyse d'impact, ecriture.

Regle d'or : rien n'est ecrit tant que appliquer() n'est pas appele explicitement.
"""

from __future__ import annotations

import datetime
import os
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional, Tuple

from . import tables
from .model import (
    Profile, ColumnRef, Injection,
    MODE_VIDES, MODE_TOUT, MODE_DIFFERENTES,
    DUP_IGNORER, DUP_PREMIER, DUP_DERNIER,
    ST_A_ECRIRE, ST_IDENTIQUE, ST_CONFLIT, ST_ORPHELIN, ST_AMBIGU, ST_SANS_CLE, ST_RIEN,
)

SEP_CLE = "\x1f"
ESPACE_INSECABLE = chr(0x00A0)   # frequent dans les exports Excel/web


# ---------------------------------------------------------------------------
#  Normalisation
# ---------------------------------------------------------------------------

def normaliser(valeur: Any, opts) -> Optional[str]:
    """Normalise une valeur de cle. Retourne None si la cellule est vide."""
    if valeur is None:
        return None
    if isinstance(valeur, bool):
        txt = "VRAI" if valeur else "FAUX"
    elif isinstance(valeur, float):
        if opts.numeric_keys and valeur.is_integer():
            txt = str(int(valeur))
        else:
            txt = repr(valeur)
    elif isinstance(valeur, int):
        txt = str(valeur)
    elif isinstance(valeur, (datetime.datetime, datetime.date)):
        txt = valeur.isoformat()
    else:
        txt = str(valeur)
        if opts.numeric_keys:
            depouille = txt.strip().replace(" ", "").replace(ESPACE_INSECABLE, "")
            if depouille and depouille.lstrip("+-").replace(".", "", 1).isdigit():
                try:
                    f = float(depouille)
                    if f.is_integer() and "." in depouille:
                        txt = str(int(f))
                except ValueError:
                    pass

    if opts.trim:
        txt = txt.strip()
    if opts.collapse_spaces:
        txt = " ".join(txt.split())
    if not txt:
        return None
    if opts.ignore_accents:
        txt = "".join(c for c in unicodedata.normalize("NFD", txt)
                      if unicodedata.category(c) != "Mn")
    if opts.ignore_separators:
        for c in (" ", "-", "_", ".", "/", "\\"):
            txt = txt.replace(c, "")
    if opts.ignore_leading_zeros:
        depouille = txt.lstrip("0")
        txt = depouille if depouille else "0"
    if opts.ignore_case:
        txt = txt.upper()
    return txt or None


def cle_laxiste(cle: Optional[str]) -> Optional[str]:
    """Variante ultra permissive : sert seulement a DETECTER les quasi-doublons."""
    if cle is None:
        return None
    txt = cle
    for c in (" ", "-", "_", ".", "/", "\\"):
        txt = txt.replace(c, "")
    txt = txt.upper().lstrip("0")
    return txt or "0"


def _nombre(valeur: Any) -> Optional[float]:
    if isinstance(valeur, bool):
        return None
    if isinstance(valeur, (int, float)):
        return float(valeur)
    if isinstance(valeur, str):
        txt = valeur.strip().replace(ESPACE_INSECABLE, "").replace(" ", "").replace(",", ".")
        if txt and txt.lstrip("+-").replace(".", "", 1).isdigit():
            try:
                return float(txt)
            except ValueError:
                return None
    return None


def memes_valeurs(a: Any, b: Any) -> bool:
    """Egalite tolerante : 12 == '12' == 12.0, espaces de bord ignores."""
    va, vb = tables.est_vide(a), tables.est_vide(b)
    if va and vb:
        return True
    if va or vb:
        return False
    na, nb = _nombre(a), _nombre(b)
    if na is not None and nb is not None:
        return abs(na - nb) < 1e-9
    return str(a).strip() == str(b).strip()


def afficher(valeur: Any) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    if isinstance(valeur, (datetime.datetime, datetime.date)):
        return valeur.isoformat(sep=" ") if isinstance(valeur, datetime.datetime) else valeur.isoformat()
    return str(valeur)


# ---------------------------------------------------------------------------
#  Resultats
# ---------------------------------------------------------------------------

@dataclass
class CellPlan:
    """Ce qui se passerait pour une injection donnee, sur une ligne donnee."""
    injection: int          # index dans profile.injections
    ancien: Any = None
    nouveau: Any = None
    action: str = ST_RIEN   # ST_A_ECRIRE / ST_IDENTIQUE / ST_CONFLIT / ST_RIEN


@dataclass
class RowResult:
    row: int
    cle: str
    statut: str
    cells: List[CellPlan] = field(default_factory=list)
    inclure: bool = True
    ligne_source: Optional[int] = None


@dataclass
class Analyse:
    lignes: List[RowResult] = field(default_factory=list)
    stats: Dict[str, int] = field(default_factory=dict)
    anomalies: List[str] = field(default_factory=list)
    avertissements: List[str] = field(default_factory=list)
    entetes_source: List[str] = field(default_factory=list)
    entetes_cible: List[str] = field(default_factory=list)
    cols_source_cle: List[int] = field(default_factory=list)
    cols_cible_cle: List[int] = field(default_factory=list)
    cols_injection: List[int] = field(default_factory=list)   # 0 = nouvelle colonne

    def a_ecrire(self) -> List[RowResult]:
        return [l for l in self.lignes if l.statut == ST_A_ECRIRE]

    def cellules_retenues(self) -> int:
        n = 0
        for l in self.lignes:
            if not l.inclure:
                continue
            n += sum(1 for c in l.cells if c.action == ST_A_ECRIRE)
        return n


class Annulation(Exception):
    """Levee quand l'operateur interrompt une analyse."""


# ---------------------------------------------------------------------------
#  Moteur
# ---------------------------------------------------------------------------

class Moteur(object):

    def __init__(self, profil: Profile):
        self.profil = profil
        self.t_source: Optional[tables.Table] = None
        self.t_cible: Optional[tables.Table] = None
        self.analyse: Optional[Analyse] = None

    # --- ouverture ---------------------------------------------------------

    def ouvrir(self) -> None:
        p = self.profil
        self.fermer()
        self.t_source = tables.ouvrir(p.source.path, p.source.sheet, p.source.header_row, writable=False)
        self.t_cible = tables.ouvrir(p.target.path, p.target.sheet, p.target.header_row, writable=True)
        # la feuille reellement ouverte peut differer (profil sans feuille memorisee)
        p.source.sheet = self.t_source.sheet
        p.target.sheet = self.t_cible.sheet

    def fermer(self) -> None:
        for t in (self.t_source, self.t_cible):
            if t is not None:
                t.close()
        self.t_source = None
        self.t_cible = None

    # --- resolution des colonnes ------------------------------------------

    @staticmethod
    def _resoudre(ref: ColumnRef, entetes: List[str], avertissements: List[str], role: str) -> int:
        """Retrouve une colonne : par en-tete d'abord, par index ensuite."""
        if ref is None:
            return 0
        idx = ref.index
        entete_a_lidx = entetes[idx - 1] if 1 <= idx <= len(entetes) else ""
        if ref.header:
            if entete_a_lidx.strip().lower() == ref.header.strip().lower():
                return idx
            correspondances = [i + 1 for i, h in enumerate(entetes)
                               if h.strip().lower() == ref.header.strip().lower()]
            if len(correspondances) == 1:
                avertissements.append(
                    "%s : la colonne '%s' a ete retrouvee en %s (le profil indiquait %s)."
                    % (role, ref.header, tables.lettre_colonne(correspondances[0]),
                       tables.lettre_colonne(idx)))
                return correspondances[0]
            if len(correspondances) > 1:
                avertissements.append(
                    "%s : l'en-tete '%s' apparait %d fois ; la colonne %s du profil est conservee."
                    % (role, ref.header, len(correspondances), tables.lettre_colonne(idx)))
            else:
                avertissements.append(
                    "%s : en-tete '%s' introuvable ; la colonne %s est utilisee telle quelle "
                    "(en-tete actuel : '%s')."
                    % (role, ref.header, tables.lettre_colonne(idx), entete_a_lidx))
        return idx

    # --- analyse -----------------------------------------------------------

    def analyser(self, progression: Optional[Callable[[str, int, int], None]] = None,
                 annule: Optional[Callable[[], bool]] = None) -> Analyse:
        p = self.profil
        if not p.keys:
            raise ValueError("Aucune cle d'association definie.")
        if not p.injections:
            raise ValueError("Aucune colonne a injecter.")
        if self.t_source is None or self.t_cible is None:
            self.ouvrir()
        if os.path.abspath(p.source.path) == os.path.abspath(p.target.path):
            raise ValueError("Le fichier source et le fichier cible sont identiques.")

        a = Analyse()
        a.entetes_source = self.t_source.headers()
        a.entetes_cible = self.t_cible.headers()

        a.cols_source_cle = [self._resoudre(k.source, a.entetes_source, a.avertissements, "Source (cle)")
                             for k in p.keys]
        a.cols_cible_cle = [self._resoudre(k.target, a.entetes_cible, a.avertissements, "Cible (cle)")
                            for k in p.keys]
        cols_src_val = [self._resoudre(i.source, a.entetes_source, a.avertissements, "Source (valeur)")
                        for i in p.injections]
        a.cols_injection = [
            self._resoudre(i.target, a.entetes_cible, a.avertissements, "Cible (destination)")
            if i.target else 0
            for i in p.injections
        ]

        self._verifier_collisions(a)

        index, ambigues = self._indexer_source(a, cols_src_val, progression, annule)
        self._analyser_cible(a, index, ambigues, progression, annule)

        self.analyse = a
        return a

    def _verifier_collisions(self, a: Analyse) -> None:
        """Signale les configurations dangereuses (ecriture dans une colonne cle...)."""
        p = self.profil
        for n, col in enumerate(a.cols_injection):
            if col and col in a.cols_cible_cle:
                a.avertissements.append(
                    "[!] L'injection %d ecrit dans %s, qui sert aussi de cle d'association."
                    % (n + 1, tables.lettre_colonne(col)))
            for m in range(n):
                if col and col == a.cols_injection[m]:
                    a.avertissements.append(
                        "[!] Les injections %d et %d visent la meme colonne %s : la derniere gagne."
                        % (m + 1, n + 1, tables.lettre_colonne(col)))

    def _indexer_source(self, a: Analyse, cols_val: List[int],
                        progression, annule) -> Tuple[Dict[str, Tuple], set]:
        p = self.profil
        t = self.t_source
        opts = p.match
        debut = t.header_row + 1
        total = max(0, t.max_row - t.header_row)

        index: Dict[str, Tuple[Any, ...]] = {}
        lignes_cle: Dict[str, List[int]] = {}
        ambigues = set()
        sans_cle: List[int] = []
        vides_par_injection = [0] * len(cols_val)
        formules: List[int] = []

        for ligne in range(debut, t.max_row + 1):
            if annule and annule():
                raise Annulation()
            if progression and (ligne - debut) % 500 == 0:
                progression("Lecture de la source", ligne - debut, total)

            parties = [normaliser(t.get(ligne, c), opts) for c in a.cols_source_cle]
            if any(x is None for x in parties):
                if any(not tables.est_vide(t.get(ligne, c)) for c in range(1, t.max_column + 1)):
                    sans_cle.append(ligne)
                continue
            cle = SEP_CLE.join(parties)
            lignes_cle.setdefault(cle, []).append(ligne)

            valeurs = []
            for n, c in enumerate(cols_val):
                v = t.get(ligne, c) if c else None
                if isinstance(v, str) and v.startswith("="):
                    formules.append(ligne)
                    v = None
                if tables.est_vide(v):
                    vides_par_injection[n] += 1
                valeurs.append(v)
            valeurs = tuple(valeurs)

            if cle not in index:
                index[cle] = valeurs
                continue
            if all(memes_valeurs(x, y) for x, y in zip(index[cle], valeurs)):
                continue
            if opts.on_duplicate == DUP_DERNIER:
                index[cle] = valeurs
            elif opts.on_duplicate == DUP_PREMIER:
                pass
            else:
                ambigues.add(cle)

        for cle in ambigues:
            index.pop(cle, None)

        doublons = {k: v for k, v in lignes_cle.items() if len(v) > 1}
        laxiste: Dict[str, List[str]] = {}
        for cle in lignes_cle:
            laxiste.setdefault(SEP_CLE.join(cle_laxiste(x) or "" for x in cle.split(SEP_CLE)), []).append(cle)
        quasi = {k: v for k, v in laxiste.items() if len(v) > 1}

        a.stats["source_lignes"] = total
        a.stats["source_cles"] = len(lignes_cle)
        a.stats["source_exploitables"] = len(index)
        a.stats["source_doublons"] = len(doublons)
        a.stats["source_ambigues"] = len(ambigues)
        a.stats["source_sans_cle"] = len(sans_cle)
        a.stats["source_quasi_doublons"] = len(quasi)

        if doublons:
            a.anomalies.append("[!] %d cle(s) en doublon dans la source :" % len(doublons))
            for cle, lignes in list(doublons.items())[:20]:
                a.anomalies.append("      %-30s lignes %s" % (cle.replace(SEP_CLE, " | "), lignes[:12]))
            if len(doublons) > 20:
                a.anomalies.append("      ... et %d autre(s)" % (len(doublons) - 20))
        if ambigues:
            a.anomalies.append("[!!] %d cle(s) en doublon avec des VALEURS DIFFERENTES (cles ecartees) :"
                               % len(ambigues))
            for cle in list(ambigues)[:20]:
                a.anomalies.append("      %s" % cle.replace(SEP_CLE, " | "))
            if len(ambigues) > 20:
                a.anomalies.append("      ... et %d autre(s)" % (len(ambigues) - 20))
        for n, nb in enumerate(vides_par_injection):
            if nb:
                a.anomalies.append("[i] Injection %d (%s) : %d ligne(s) source sans valeur."
                                   % (n + 1, p.injections[n].source.label(), nb))
        if sans_cle:
            a.anomalies.append("[!] %d ligne(s) source non vides mais sans cle : %s"
                               % (len(sans_cle), sans_cle[:15]))
        if formules:
            a.anomalies.append("[!!] %d cellule(s) source contiennent une formule sans valeur en cache "
                               "(lignes %s...). Ouvrez le fichier dans Excel, enregistrez, relancez."
                               % (len(formules), sorted(set(formules))[:10]))
        if quasi:
            a.anomalies.append("[!] %d groupe(s) de cles tres proches (zeros de tete / separateurs), "
                               "traitees comme differentes :" % len(quasi))
            for _, variantes in list(quasi.items())[:10]:
                a.anomalies.append("      %s" % [v.replace(SEP_CLE, " | ") for v in variantes])

        return index, ambigues

    def _analyser_cible(self, a: Analyse, index: Dict[str, Tuple], ambigues: set,
                        progression, annule) -> None:
        p = self.profil
        t = self.t_cible
        opts = p.match
        debut = t.header_row + 1
        total = max(0, t.max_row - t.header_row)

        compteurs = {k: 0 for k in
                     (ST_A_ECRIRE, ST_IDENTIQUE, ST_CONFLIT, ST_ORPHELIN, ST_AMBIGU, ST_SANS_CLE, ST_RIEN)}
        cellules = 0
        cles_vues = set()

        for ligne in range(debut, t.max_row + 1):
            if annule and annule():
                raise Annulation()
            if progression and (ligne - debut) % 500 == 0:
                progression("Analyse de la cible", ligne - debut, total)

            parties = [normaliser(t.get(ligne, c), opts) for c in a.cols_cible_cle]
            if any(x is None for x in parties):
                if any(not tables.est_vide(t.get(ligne, c)) for c in range(1, t.max_column + 1)):
                    compteurs[ST_SANS_CLE] += 1
                    a.lignes.append(RowResult(ligne, "", ST_SANS_CLE, []))
                continue

            cle = SEP_CLE.join(parties)
            affiche = cle.replace(SEP_CLE, " | ")
            cles_vues.add(cle)

            if cle in ambigues:
                compteurs[ST_AMBIGU] += 1
                a.lignes.append(RowResult(ligne, affiche, ST_AMBIGU, []))
                continue

            trouve = cle in index
            valeurs = index.get(cle)

            if not trouve and not any(i.apply_default for i in p.injections):
                compteurs[ST_ORPHELIN] += 1
                a.lignes.append(RowResult(ligne, affiche, ST_ORPHELIN, []))
                continue

            cells: List[CellPlan] = []
            for n, inj in enumerate(p.injections):
                col = a.cols_injection[n]
                ancien = t.get(ligne, col) if col else None
                if trouve:
                    nouveau = valeurs[n]
                elif inj.apply_default:
                    nouveau = inj.default
                else:
                    cells.append(CellPlan(n, ancien, None, ST_RIEN))
                    continue

                if tables.est_vide(nouveau):
                    cells.append(CellPlan(n, ancien, nouveau, ST_RIEN))
                    continue
                if memes_valeurs(ancien, nouveau):
                    cells.append(CellPlan(n, ancien, nouveau, ST_IDENTIQUE))
                    continue
                if tables.est_vide(ancien):
                    cells.append(CellPlan(n, ancien, nouveau, ST_A_ECRIRE))
                    continue
                # cellule deja remplie avec autre chose
                if inj.mode in (MODE_TOUT, MODE_DIFFERENTES):
                    cells.append(CellPlan(n, ancien, nouveau, ST_A_ECRIRE))
                else:
                    cells.append(CellPlan(n, ancien, nouveau, ST_CONFLIT))

            actions = [c.action for c in cells]
            if ST_A_ECRIRE in actions:
                statut = ST_A_ECRIRE
                cellules += actions.count(ST_A_ECRIRE)
            elif ST_CONFLIT in actions:
                statut = ST_CONFLIT
            elif ST_IDENTIQUE in actions:
                statut = ST_IDENTIQUE
            elif not trouve:
                statut = ST_ORPHELIN
            else:
                statut = ST_RIEN
            compteurs[statut] += 1
            a.lignes.append(RowResult(ligne, affiche, statut, cells))

        non_utilisees = sorted(set(index) - cles_vues)

        a.stats["cible_lignes"] = total
        a.stats["cible_cles"] = len(cles_vues)
        a.stats["cellules_a_ecrire"] = cellules
        for k, v in compteurs.items():
            a.stats["statut_" + k] = v
        a.stats["source_non_utilisees"] = len(non_utilisees)

        if non_utilisees:
            a.anomalies.append("[i] %d cle(s) de la source sans equivalent dans la cible : %s%s"
                               % (len(non_utilisees),
                                  [c.replace(SEP_CLE, " | ") for c in non_utilisees[:20]],
                                  " ..." if len(non_utilisees) > 20 else ""))

    # --- ecriture ----------------------------------------------------------

    def appliquer(self, progression: Optional[Callable[[str, int, int], None]] = None) -> Dict[str, Any]:
        """Ecrit les cellules retenues. Retourne un dict de comptes + chemins."""
        p = self.profil
        a = self.analyse
        if a is None:
            raise ValueError("Lancez d'abord l'analyse.")
        t = self.t_cible

        # 1) creation des colonnes neuves
        for n, inj in enumerate(p.injections):
            if not a.cols_injection[n]:
                entete = inj.new_header or ("%s (INJXL)" % (inj.source.header or inj.source.label()))
                a.cols_injection[n] = t.ajouter_colonne(entete)

        # 2) ecriture
        retenues = [l for l in a.lignes if l.inclure and l.statut == ST_A_ECRIRE]
        ecrites = 0
        journal: List[Tuple[int, str, str, str, str]] = []
        for i, res in enumerate(retenues):
            if progression and i % 200 == 0:
                progression("Ecriture", i, len(retenues))
            for cell in res.cells:
                if cell.action != ST_A_ECRIRE:
                    continue
                col = a.cols_injection[cell.injection]
                t.set(res.row, col, cell.nouveau)
                if p.output.highlight:
                    try:
                        t.surligner(res.row, col, p.output.highlight_color)
                    except Exception:
                        pass
                ecrites += 1
                journal.append((res.row, res.cle, tables.lettre_colonne(col),
                                afficher(cell.ancien), afficher(cell.nouveau)))

        # 3) enregistrement
        sauvegarde = ""
        if p.output.mode == "surplace":
            destination = os.path.abspath(p.target.path)
            if p.output.backup:
                sauvegarde = tables.sauvegarder_copie(destination)
        else:
            destination = p.output.path.strip() or tables.chemin_auto(p.target.path)
            if os.path.isdir(destination):
                destination = os.path.join(
                    destination, os.path.basename(tables.chemin_auto(p.target.path)))
        dossier = os.path.dirname(os.path.abspath(destination))
        if dossier and not os.path.isdir(dossier):
            os.makedirs(dossier)
        t.save(destination)

        return {
            "cellules": ecrites,
            "lignes": len(retenues),
            "destination": destination,
            "sauvegarde": sauvegarde,
            "journal": journal,
        }

    # --- rapports ----------------------------------------------------------

    def rapport_texte(self, resultat: Optional[Dict[str, Any]] = None) -> str:
        p = self.profil
        a = self.analyse
        L: List[str] = []
        trait = "=" * 78
        L.append(trait)
        L.append("RAPPORT INJXL - %s" % datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
        L.append(trait)
        L.append("Profil            : %s" % p.name)
        L.append("Fichier source    : %s" % os.path.abspath(p.source.path))
        L.append("  feuille         : %s   (en-tete ligne %d)" % (p.source.sheet, p.source.header_row))
        L.append("Fichier cible     : %s" % os.path.abspath(p.target.path))
        L.append("  feuille         : %s   (en-tete ligne %d)" % (p.target.sheet, p.target.header_row))
        L.append("")
        L.append("Cle d'association :")
        for n, k in enumerate(p.keys):
            L.append("  %d. source %-24s  <->  cible %s"
                     % (n + 1, tables.lettre_colonne(a.cols_source_cle[n]) + " " + (k.source.header or ""),
                        tables.lettre_colonne(a.cols_cible_cle[n]) + " " + (k.target.header or "")))
        L.append("")
        L.append("Injections :")
        from .model import MODES_LIBELLES
        for n, inj in enumerate(p.injections):
            col = a.cols_injection[n]
            cible = tables.lettre_colonne(col) if col else "[nouvelle colonne]"
            L.append("  %d. %s  ->  %s   (%s)"
                     % (n + 1, inj.source.label(), cible, MODES_LIBELLES.get(inj.mode, inj.mode)))
        L.append("")
        L.append("Options de comparaison : casse=%s espaces=%s accents=%s separateurs=%s "
                 "zeros=%s doublons=%s"
                 % (not p.match.ignore_case, not p.match.collapse_spaces,
                    not p.match.ignore_accents, not p.match.ignore_separators,
                    not p.match.ignore_leading_zeros, p.match.on_duplicate))
        L.append("")
        L.append("-" * 78)
        L.append("1) SOURCE")
        L.append("-" * 78)
        s = a.stats
        L.append("  Lignes de donnees             : %s" % s.get("source_lignes", 0))
        L.append("  Cles uniques                  : %s" % s.get("source_cles", 0))
        L.append("  Cles exploitables             : %s" % s.get("source_exploitables", 0))
        L.append("  Cles en doublon               : %s" % s.get("source_doublons", 0))
        L.append("  Cles ecartees (ambigues)      : %s" % s.get("source_ambigues", 0))
        L.append("")
        L.append("-" * 78)
        L.append("2) IMPACT SUR LA CIBLE")
        L.append("-" * 78)
        L.append("  Lignes de donnees             : %s" % s.get("cible_lignes", 0))
        L.append("  Cles uniques                  : %s" % s.get("cible_cles", 0))
        L.append("  Lignes a mettre a jour        : %s" % s.get("statut_" + ST_A_ECRIRE, 0))
        L.append("  CELLULES A ECRIRE             : %s" % s.get("cellules_a_ecrire", 0))
        L.append("  Deja identiques               : %s" % s.get("statut_" + ST_IDENTIQUE, 0))
        L.append("  Conflits conserves            : %s" % s.get("statut_" + ST_CONFLIT, 0))
        L.append("  Orphelins (absents de source) : %s" % s.get("statut_" + ST_ORPHELIN, 0))
        L.append("  Lignes sans cle               : %s" % s.get("statut_" + ST_SANS_CLE, 0))
        L.append("  Cles source non utilisees     : %s" % s.get("source_non_utilisees", 0))
        if a.avertissements:
            L.append("")
            L.append("-" * 78)
            L.append("3) AVERTISSEMENTS DE CONFIGURATION")
            L.append("-" * 78)
            L.extend("  " + x for x in a.avertissements)
        if a.anomalies:
            L.append("")
            L.append("-" * 78)
            L.append("4) ANOMALIES DETECTEES")
            L.append("-" * 78)
            L.extend("  " + x for x in a.anomalies)

        conflits = [l for l in a.lignes if l.statut == ST_CONFLIT]
        if conflits:
            L.append("")
            L.append("-" * 78)
            L.append("5) DETAIL DES CONFLITS (valeurs cibles conservees)")
            L.append("-" * 78)
            for l in conflits[:100]:
                for c in l.cells:
                    if c.action == ST_CONFLIT:
                        L.append("  ligne %-7d %-28s actuel=%r  source=%r"
                                 % (l.row, l.cle, afficher(c.ancien), afficher(c.nouveau)))
            if len(conflits) > 100:
                L.append("  ... et %d ligne(s) supplementaire(s)" % (len(conflits) - 100))

        if resultat:
            L.append("")
            L.append("-" * 78)
            L.append("6) ECRITURE EFFECTUEE")
            L.append("-" * 78)
            L.append("  Cellules ecrites   : %s" % resultat["cellules"])
            L.append("  Lignes modifiees   : %s" % resultat["lignes"])
            L.append("  Fichier produit    : %s" % resultat["destination"])
            if resultat.get("sauvegarde"):
                L.append("  Sauvegarde         : %s" % resultat["sauvegarde"])
            L.append("")
            L.append("  %-8s %-28s %-4s %-20s %s" % ("Ligne", "Cle", "Col", "Ancien", "Nouveau"))
            for row, cle, col, ancien, nouveau in resultat["journal"]:
                L.append("  %-8d %-28s %-4s %-20s %s" % (row, cle[:28], col, ancien[:20], nouveau))
        else:
            L.append("")
            L.append("  (simulation - aucune donnee ecrite)")
        L.append("")
        return "\n".join(L)

    def ecrire_rapport_txt(self, texte: str, base: str) -> str:
        chemin = "%s_RAPPORT_%s.txt" % (os.path.splitext(base)[0], tables.horodatage())
        with open(chemin, "w", encoding="utf-8-sig") as f:
            f.write(texte)
        return chemin

    def ecrire_rapport_xlsx(self, base: str) -> str:
        """Rapport tabulaire : une feuille de synthese + le detail ligne a ligne."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        a = self.analyse
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Synthese"
        gras = Font(bold=True)
        ws.append(["Indicateur", "Valeur"])
        ws["A1"].font = gras
        ws["B1"].font = gras
        libelles = [
            ("Lignes source", "source_lignes"), ("Cles source uniques", "source_cles"),
            ("Cles exploitables", "source_exploitables"), ("Cles en doublon", "source_doublons"),
            ("Cles ecartees", "source_ambigues"), ("Lignes cible", "cible_lignes"),
            ("Cles cible uniques", "cible_cles"), ("Cellules a ecrire", "cellules_a_ecrire"),
            ("Lignes a ecrire", "statut_" + ST_A_ECRIRE), ("Identiques", "statut_" + ST_IDENTIQUE),
            ("Conflits", "statut_" + ST_CONFLIT), ("Orphelins", "statut_" + ST_ORPHELIN),
            ("Sans cle", "statut_" + ST_SANS_CLE), ("Cles source inutilisees", "source_non_utilisees"),
        ]
        for libelle, cle in libelles:
            ws.append([libelle, a.stats.get(cle, 0)])
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 14

        wd = wb.create_sheet("Detail")
        entetes = ["Ligne", "Cle", "Statut"]
        for n, inj in enumerate(self.profil.injections):
            entetes += ["Inj%d actuel" % (n + 1), "Inj%d propose" % (n + 1), "Inj%d action" % (n + 1)]
        wd.append(entetes)
        for c in wd[1]:
            c.font = gras
        couleurs = {
            ST_A_ECRIRE: "FFC6EFCE", ST_CONFLIT: "FFFFC7CE",
            ST_ORPHELIN: "FFFFEB9C", ST_AMBIGU: "FFFFEB9C",
        }
        for l in a.lignes:
            ligne = [l.row, l.cle, l.statut]
            for n in range(len(self.profil.injections)):
                cell = next((c for c in l.cells if c.injection == n), None)
                ligne += [afficher(cell.ancien) if cell else "",
                          afficher(cell.nouveau) if cell else "",
                          cell.action if cell else ""]
            wd.append(ligne)
            if l.statut in couleurs:
                remplissage = PatternFill(start_color=couleurs[l.statut],
                                          end_color=couleurs[l.statut], fill_type="solid")
                wd.cell(row=wd.max_row, column=3).fill = remplissage
        wd.freeze_panes = "A2"
        wd.auto_filter.ref = wd.dimensions
        wd.column_dimensions["B"].width = 28
        wd.column_dimensions["C"].width = 15

        wa = wb.create_sheet("Anomalies")
        wa.append(["Message"])
        wa["A1"].font = gras
        for x in a.avertissements + a.anomalies:
            wa.append([x])
        wa.column_dimensions["A"].width = 120

        chemin = "%s_RAPPORT_%s.xlsx" % (os.path.splitext(base)[0], tables.horodatage())
        wb.save(chemin)
        return chemin
